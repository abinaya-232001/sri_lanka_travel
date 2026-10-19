"""Load police stations, tourism police and hospitals, and create one login per station.

    python scripts/seed.py              # first time
    python scripts/seed.py --demo       # also create a demo tourist account

Officer usernames look like "ella-police". Every station gets a RANDOM password,
written once to data/officer_credentials.csv. Give each station its own line,
then delete the file. Officers must change the password at first login.
"""

import argparse
import csv
import re
import secrets
import string
import sys
from pathlib import Path

from openpyxl import load_workbook

BACKEND = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND))

from app import create_app  # noqa: E402
from app.extensions import db  # noqa: E402
from app.models import Hospital, Officer, Station, Tourist  # noqa: E402

DATA = BACKEND / "data"


def rows(path):
    ws = load_workbook(path, read_only=True, data_only=True).active
    it = ws.iter_rows(values_only=True)
    head = [str(h).strip() if h is not None else "" for h in next(it)]
    for r in it:
        yield dict(zip(head, r))


def phone(v):
    if v is None:
        return None
    s = re.sub(r"\D", "", str(v).split(".")[0])
    if len(s) == 9:
        s = "0" + s  # numbers stored as integers lost their leading 0
    return s or None


def slug(name):
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")


def password():
    alphabet = string.ascii_letters + string.digits
    while True:
        p = "".join(secrets.choice(alphabet) for _ in range(12))
        if any(c.isdigit() for c in p) and any(c.isalpha() for c in p):
            return p


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--demo", action="store_true", help="create demo tourist demo@tourceylon.lk")
    args = ap.parse_args()
    app = create_app()
    with app.app_context():
        existing = {s.name for s in Station.query.all()}
        new_stations = []
        for r in rows(DATA / "police_stations.xlsx"):
            if r.get("Latitude") is None or not r.get("Division"):
                continue
            name = f"{str(r['Division']).strip()} Police"
            if name not in existing:
                new_stations.append(Station(name=name, kind=Station.KIND_POLICE, telephone=phone(r.get("Telephone No")),
                                            latitude=float(r["Latitude"]), longitude=float(r["Longitude"])))
                existing.add(name)
        for r in rows(DATA / "tourism_police_stations.xlsx"):
            if r.get("Latitude") is None or not r.get("Police Station"):
                continue
            name = f"{str(r['Police Station']).strip()} Tourism Police"
            if name not in existing:
                new_stations.append(Station(name=name, kind=Station.KIND_TOURISM,
                                            telephone=phone(r.get("Phone 1") or r.get("Phone 2")),
                                            latitude=float(r["Latitude"]), longitude=float(r["Longitude"])))
                existing.add(name)
        db.session.add_all(new_stations)

        if not Hospital.query.first():
            for r in rows(DATA / "hospitals.xlsx"):
                if r.get("Latitude") is None or not r.get("Hospital"):
                    continue
                db.session.add(Hospital(name=str(r["Hospital"]).strip(), district=r.get("District"),
                                        phone=phone(r.get("Local Number")),
                                        latitude=float(r["Latitude"]), longitude=float(r["Longitude"])))
        db.session.commit()

        creds, used = [], {o.username for o in Officer.query.all()}
        for st in Station.query.order_by(Station.id).all():
            if Officer.query.filter_by(station_id=st.id).first():
                continue
            base = uname = slug(st.name)
            n = 2
            while uname in used:
                uname, n = f"{base}-{n}", n + 1
            used.add(uname)
            pw = password()
            o = Officer(username=uname, station_id=st.id, must_change_password=True)
            o.set_password(pw)
            db.session.add(o)
            creds.append((st.name, st.kind, uname, pw))
        db.session.commit()

        if creds:
            out = DATA / "officer_credentials.csv"
            with open(out, "w", newline="", encoding="utf-8") as fh:
                w = csv.writer(fh)
                w.writerow(["station", "kind", "username", "temporary_password"])
                w.writerows(creds)
            print(f"Created {len(creds)} officer logins -> {out}  (share securely, then delete the file)")

        if args.demo and not Tourist.query.filter_by(email="demo@tourceylon.lk").first():
            t = Tourist(name="Demo Tourist", email="demo@tourceylon.lk", phone="+447700900123",
                        nationality="British", language="en")
            t.set_password("demo12345")
            db.session.add(t)
            db.session.commit()
            print("Demo tourist: demo@tourceylon.lk / demo12345")

        print(f"Stations: {Station.query.count()}  Hospitals: {Hospital.query.count()}  Officers: {Officer.query.count()}")


if __name__ == "__main__":
    main()
